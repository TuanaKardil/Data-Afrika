"""Resumable Excel writer: streams rows to a staging CSV, then converts to 4-sheet XLSX."""

from __future__ import annotations

import csv
import io
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import AuditEntry, EnrichedRow


class StagingWriter:
    """Appends enriched rows to a CSV file after each company is processed."""

    def __init__(self, staging_path: Path) -> None:
        self._path = staging_path
        self._file: io.TextIOWrapper | None = None
        self._writer: csv.DictWriter[str] | None = None
        self._fieldnames: list[str] | None = None

    def _ensure_open(self, row: EnrichedRow) -> None:
        if self._writer is not None:
            return
        fieldnames = list(type(row).model_fields.keys())
        self._fieldnames = fieldnames
        exists = self._path.exists()
        self._file = open(self._path, "a", newline="", encoding="utf-8")  # noqa: SIM115
        self._writer = csv.DictWriter(self._file, fieldnames=fieldnames, extrasaction="ignore")
        if not exists:
            self._writer.writeheader()

    def write_row(self, row: EnrichedRow) -> None:
        self._ensure_open(row)
        assert self._writer is not None
        data = {k: v for k, v in row.model_dump().items()}
        for k, v in list(data.items()):
            if isinstance(v, dict):
                data[k] = str(v)
        self._writer.writerow(data)
        assert self._file is not None
        self._file.flush()

    def close(self) -> None:
        if self._file:
            self._file.close()


def load_staging(staging_path: Path) -> set[int]:
    """Return set of row_index values already present in the staging CSV."""
    if not staging_path.exists():
        return set()
    try:
        df = pd.read_csv(staging_path, usecols=["row_index"], dtype={"row_index": int})
        return set(df["row_index"].tolist())
    except Exception:
        return set()


# Columns from EnrichedRow that are internal/technical — excluded from the clean results sheet
_INTERNAL_COLS = {
    "row_index", "normalized_name", "country_iso2", "inferred_sector",
    "cleaned_name", "extra", "audit_ids", "processed_at",
    "gps_lat", "gps_lon", "address", "instagram_url", "whatsapp_number",
    # per-phone/email detail columns — kept only as number+source
    *(f"phone_{i}_{s}" for i in range(1, 6) for s in ("confidence", "type", "audit_id")),
    *(f"email_{i}_{s}" for i in range(1, 3) for s in ("confidence", "validation", "audit_id")),
}

_HEADER_LABELS: dict[str, str] = {
    "importer_name": "Firma",
    "country": "Ülke",
    "phone_1": "Telefon 1", "phone_1_source": "Tel1 Kaynak",
    "phone_2": "Telefon 2", "phone_2_source": "Tel2 Kaynak",
    "phone_3": "Telefon 3", "phone_3_source": "Tel3 Kaynak",
    "phone_4": "Telefon 4", "phone_4_source": "Tel4 Kaynak",
    "phone_5": "Telefon 5", "phone_5_source": "Tel5 Kaynak",
    "email_1": "E-posta 1", "email_1_source": "Email1 Kaynak",
    "email_2": "E-posta 2", "email_2_source": "Email2 Kaynak",
    "website": "Website",
    "facebook_url": "Facebook",
    "linkedin_url": "LinkedIn",
    "search_status": "Durum",
    "overall_confidence": "Güven (%)",
    "notes": "Notlar",
}

_COLUMN_ORDER = [
    "importer_name", "country",
    "phone_1", "phone_1_source",
    "phone_2", "phone_2_source",
    "phone_3", "phone_3_source",
    "phone_4", "phone_4_source",
    "phone_5", "phone_5_source",
    "email_1", "email_1_source",
    "email_2", "email_2_source",
    "website", "facebook_url", "linkedin_url",
    "search_status", "overall_confidence", "notes",
]


_PHONE_COLS = [f"phone_{i}" for i in range(1, 6)]
_EMAIL_COLS = [f"email_{i}" for i in range(1, 3)]


def _remove_cross_row_duplicates(df: pd.DataFrame) -> pd.DataFrame:
    """
    Post-processing pass: blank out phones/emails that appear in 2+ different rows.
    These are broker/intermediary contacts that must not be attributed to any single company.
    """
    if df.empty:
        return df

    # Count how many rows each phone/email appears in
    phone_counter: Counter[str] = Counter()
    email_counter: Counter[str] = Counter()

    for col in _PHONE_COLS:
        if col not in df.columns:
            continue
        for val in df[col].dropna():
            v = str(val).strip()
            if v and v not in ("nan", ""):
                phone_counter[v] += 1

    for col in _EMAIL_COLS:
        if col not in df.columns:
            continue
        for val in df[col].dropna():
            v = str(val).strip().lower()
            if v and v not in ("nan", ""):
                email_counter[v] += 1

    shared_phones = {p for p, c in phone_counter.items() if c >= 2}
    shared_emails = {e for e, c in email_counter.items() if c >= 2}

    if not shared_phones and not shared_emails:
        return df

    df = df.copy()

    def _blank_row(row: Any) -> Any:
        extra_notes: list[str] = []

        for col in _PHONE_COLS:
            if col not in row.index:
                continue
            val = str(row[col]).strip() if pd.notna(row[col]) else ""
            if val in shared_phones:
                row[col] = pd.NA
                src = col + "_source"
                if src in row.index:
                    row[src] = pd.NA
                extra_notes.append(f"Filtered: phone {val} shared across companies")

        for col in _EMAIL_COLS:
            if col not in row.index:
                continue
            val = str(row[col]).strip().lower() if pd.notna(row[col]) else ""
            if val in shared_emails:
                row[col] = pd.NA
                src = col + "_source"
                if src in row.index:
                    row[src] = pd.NA
                extra_notes.append(f"Filtered: email {val} shared across companies")

        if extra_notes:
            existing = str(row.get("notes", "") or "")
            combined = "; ".join(filter(None, [existing] + extra_notes))
            row["notes"] = combined

        return row

    return df.apply(_blank_row, axis=1)  # type: ignore[no-any-return]


def _clean_results(df: pd.DataFrame) -> pd.DataFrame:
    """Produce a clean, human-readable results dataframe."""
    all_model_cols = set(EnrichedRow.model_fields.keys())

    # Ordered standard columns that exist in df
    ordered = [c for c in _COLUMN_ORDER if c in df.columns]

    # Extra input columns not part of the model (passthrough from original xlsx)
    extras = [c for c in df.columns if c not in all_model_cols and c not in ordered]
    cols = ordered + extras

    clean = df[cols].copy()

    # Force phone columns to string to prevent scientific notation in Excel
    phone_cols = [c for c in clean.columns if c.startswith("phone_") and "_source" not in c]
    for col in phone_cols:
        clean[col] = clean[col].apply(
            lambda v: str(int(float(v))) if pd.notna(v) and str(v) not in ("", "nan") else pd.NA
        )

    # Replace empty strings / "nan" strings with NaN so Excel cells are blank
    clean.replace({"": pd.NA, "nan": pd.NA}, inplace=True)

    # Rename to human-readable headers
    clean.rename(columns=_HEADER_LABELS, inplace=True)

    return clean


def _format_sheet(ws: Any, freeze_row: int = 1) -> None:
    """Apply header formatting, freeze pane, and auto column widths."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    body_font = Font(size=10)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    ws.freeze_panes = ws.cell(row=freeze_row + 1, column=1)

    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
            else:
                cell.font = body_font
        col_letter = get_column_letter(col_idx)
        # Cap width: min 10, max 40
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    ws.row_dimensions[1].height = 22


def finalize(
    staging_path: Path,
    output_path: Path,
    audit_log: list[AuditEntry],
    errors: list[dict[str, Any]],
) -> None:
    """Convert staging CSV + audit data into a 4-sheet XLSX file."""
    raw_df = pd.read_csv(staging_path) if staging_path.exists() else pd.DataFrame()
    if not raw_df.empty:
        raw_df = _remove_cross_row_duplicates(raw_df)
    results_df = _clean_results(raw_df) if not raw_df.empty else pd.DataFrame()

    audit_records = [e.model_dump() for e in audit_log]
    audit_df = pd.DataFrame(audit_records) if audit_records else pd.DataFrame()

    if audit_records:
        adf = pd.DataFrame(audit_records)
        stats = (
            adf.groupby("source_name")
            .agg(
                attempts=("status", "count"),
                hits=("status", lambda x: (x == "hit").sum()),
                misses=("status", lambda x: (x == "miss").sum()),
                blocks=("status", lambda x: (x == "blocked").sum()),
                avg_latency_ms=("latency_ms", "mean"),
            )
            .reset_index()
        )
        stats["hit_rate_pct"] = (stats["hits"] / stats["attempts"] * 100).round(1)
        stats_df = stats
    else:
        stats_df = pd.DataFrame()

    errors_df = pd.DataFrame(errors) if errors else pd.DataFrame()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        results_df.to_excel(writer, sheet_name="Sonuçlar", index=False)
        audit_df.to_excel(writer, sheet_name="Audit", index=False)
        stats_df.to_excel(writer, sheet_name="Kaynak İstatistik", index=False)
        errors_df.to_excel(writer, sheet_name="Hatalar", index=False)

        wb = writer.book
        for sheet_name in wb.sheetnames:
            _format_sheet(wb[sheet_name])
